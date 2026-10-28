# -*- coding: utf-8 -*-
"""
สร้างไฟล์ Excel รวมข้อมูลบอททั้งหมด: reports/steady_stack.xlsx
เรียกจาก dca_bot.py (ทุกชั่วโมง) และ price_watch.py (ทุกครั้งที่มีการซื้อ/ขาย/เติม/ถอน)
แท็บ: ภาพรวม · พอร์ตตอนนี้ · ระบบตัดสินใจ · ประวัติซื้อ · ประวัติขาย · กำไรขาดทุน · มูลค่าพอร์ต · เหตุการณ์
อ่านอย่างเดียวจากไฟล์ state/analysis/latest.json — ถ้าสร้างไม่สำเร็จ ไม่กระทบการเทรด
"""

import glob
import json
import os
from collections import OrderedDict, defaultdict
from datetime import datetime, timedelta, timezone

REPORT_FILE = "reports/steady_stack.xlsx"
BKK = timezone(timedelta(hours=7))
FONT = "Arial"
# ทุนของบอทก่อนเปลี่ยนเป็นโหมดทั้งบัญชี (เติม 400 + 300 บาท ตามที่ผู้ใช้แจ้ง) ใช้คิดกำไรเทียบทุนช่วงก่อน 21 ต.ค. 2569
BOT_CAPITAL_BEFORE_WHOLE_THB = 700.0


def _load(path, default):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return default


def _jsonl(path):
    out = []
    try:
        with open(path, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        out.append(json.loads(line))
                    except ValueError:
                        pass
    except OSError:
        pass
    return out


def _bkk(ts):
    """ISO (UTC) หรือ 'YYYY-MM-DD' -> datetime เวลาไทย (ไม่มี tz ให้ Excel แสดงตรงๆ)"""
    if not ts:
        return None
    try:
        if len(ts) == 10:
            return datetime.strptime(ts, "%Y-%m-%d")
        dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(BKK).replace(tzinfo=None)
    except ValueError:
        return None


def _capital_of(e):
    """ทุนของแต่ละจุดเวลา: โหมดทั้งบัญชีใช้ capital_thb · ก่อนหน้านั้น = เงินที่ใส่เข้าบอท
    (ทุนเหรียญที่ถือ + เงินในกอง + กำไรแยกเก็บ - กำไร/ขาดทุนที่ขายแล้วสะสม)"""
    if e.get("capital_thb"):
        return e["capital_thb"]
    try:
        return round((e.get("holdings_invested_thb") or 0) + (e.get("cash_pool_thb") or 0)
                     + (e.get("profit_reserved_thb") or 0) - (e.get("realized_pnl_total_thb") or 0), 2)
    except TypeError:
        return BOT_CAPITAL_BEFORE_WHOLE_THB


def build(dca_module=None, path=REPORT_FILE):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    m = dca_module
    state = _load("state/dca_state.json", {})
    journal = _load("state/trade_journal.json", {"lots": []})
    latest = _load("latest.json", {})
    dca = latest.get("dca_bot") or {}

    H = Font(name=FONT, bold=True, color="FFFFFF")
    HF = PatternFill("solid", fgColor="1F4E78")
    N = Font(name=FONT)
    B = Font(name=FONT, bold=True)
    I = Font(name=FONT, italic=True, color="595959")
    T = Font(name=FONT, bold=True, size=14)
    GOOD = Font(name=FONT, color="008000")
    BAD = Font(name=FONT, color="C00000")
    thin = Side(style="thin", color="D9D9D9")
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    WRAP = Alignment(wrap_text=True, vertical="top")
    DT = "yyyy-mm-dd hh:mm"
    THB = "#,##0.00;-#,##0.00"
    PCT = "+0.00%;-0.00%;0.00%"

    def title(ws, text, sub=None):
        ws["A1"] = text
        ws["A1"].font = T
        if sub:
            ws["A2"] = sub
            ws["A2"].font = I

    def table(ws, row, headers, rows, widths=None, formats=None):
        for c, h in enumerate(headers, 1):
            x = ws.cell(row=row, column=c, value=h)
            x.font, x.fill, x.alignment, x.border = H, HF, WRAP, BOX
        for r, vals in enumerate(rows, row + 1):
            for c, v in enumerate(vals, 1):
                x = ws.cell(row=r, column=c, value=v)
                x.font, x.border = N, BOX
                fmt = (formats or {}).get(c)
                if fmt:
                    x.number_format = fmt
                elif isinstance(v, datetime):
                    x.number_format = DT
        if widths:
            for i, w in enumerate(widths):
                ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = ws.cell(row=row + 1, column=1)
        return row + len(rows)

    now = datetime.now(BKK).replace(tzinfo=None)
    wb = Workbook()

    # ---------------- ภาพรวม ----------------
    ws = wb.active
    ws.title = "ภาพรวม"
    title(ws, "🧱 Steady Stack — ภาพรวม", f"อัปเดตอัตโนมัติ {now:%Y-%m-%d %H:%M} (เวลาไทย) · ตัวเลขเป็นบาท")
    holdings = {n: h for n, h in (state.get("holdings") or {}).items() if (h or {}).get("total_coin", 0) > 0}
    value_now = 0.0
    for n, h in holdings.items():
        px = (latest.get(n) or {}).get("price")
        if px:
            value_now += h["total_coin"] * px
    dips = state.get("dip_positions") or {}
    for n, d in dips.items():
        px = (latest.get(n) or {}).get("price") or d.get("current_price")
        if px:
            value_now += d.get("coin", 0.0) * px
    pool = state.get("cash_pool_thb", 0.0)
    capital = state.get("capital_thb")
    closed = [l for l in journal.get("lots", []) if l.get("status") == "closed" and l.get("pnl_pct") is not None]
    wins = [l for l in closed if l.get("win")]
    sells = state.get("sell_log") or []
    rows = [
        ("ทุนทั้งบัญชี (ตั้งต้น + เติม)", capital, THB),
        ("มูลค่าเหรียญตอนนี้", round(value_now, 2), THB),
        ("เงินบาทในกอง", pool, THB),
        ("มูลค่าทั้งบัญชี", "=B5+B6", THB),
        ("กำไร/ขาดทุนเทียบทุน", "=IF(B4=\"\",\"\",B7-B4)", THB),
        ("กำไร/ขาดทุน %", "=IF(OR(B4=\"\",B4=0),\"\",B8/B4)", PCT),
        ("กำไรที่ขายแล้วสะสม (ตั้งแต่เริ่มบอท)", state.get("realized_pnl_total_thb", 0.0), THB),
        ("ถอนออกสะสม", state.get("withdrawn_total_thb", 0.0), THB),
        ("แจ้งเตือนถอนเมื่อกำไรเทียบทุนถึง", getattr(m, "PROFIT_WITHDRAW_ALERT_THB", 300), THB),
        ("ไม้ที่ปิดแล้ว (journal)", len(closed), "0"),
        ("ชนะ", len(wins), "0"),
        ("Winrate", (len(wins) / len(closed)) if closed else None, "0.0%"),
        ("กำไรเฉลี่ยต่อไม้", (sum(l["pnl_pct"] for l in closed) / len(closed) / 100) if closed else None, PCT),
        ("จำนวนครั้งที่ขาย (ทั้งหมด)", len(sells), "0"),
        ("เหรียญที่ถืออยู่", len(holdings), "0"),
        ("โหมดทั้งบัญชีเริ่มเมื่อ", _bkk(state.get("whole_account_since")), DT),
    ]
    ws["A3"], ws["B3"] = "รายการ", "ค่า"
    for c in ("A3", "B3"):
        ws[c].font, ws[c].fill = H, HF
    for i, (k, v, fmt) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k).font = N
        x = ws.cell(row=i, column=2, value=v)
        x.font, x.number_format = B, fmt
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    ws.cell(row=21, column=1, value="กราฟมูลค่าพอร์ตอยู่ในแท็บ \"มูลค่าพอร์ต\" · กำไรรายเหรียญอยู่ในแท็บ \"กำไรขาดทุน\"").font = I

    # ---------------- พอร์ตตอนนี้ ----------------
    ws = wb.create_sheet("พอร์ตตอนนี้")
    title(ws, "เหรียญที่ถืออยู่", f"ราคาจากผลวิเคราะห์ล่าสุด {(_bkk(latest.get('updated_at')) or now):%Y-%m-%d %H:%M} · ราคาขายคำนวณจากทุนเฉลี่ย")
    tp = getattr(m, "TAKE_PROFIT_PCT", 6.0) / 100
    sl = getattr(m, "STOP_LOSS_PCT", 20.0) / 100
    cz = getattr(m, "CUT_ZONE_PCT", 10.0) / 100
    prow = []
    for n, h in sorted(holdings.items(), key=lambda kv: -kv[1]["total_invested_thb"]):
        info = latest.get(n) or {}
        avg = h["total_invested_thb"] / h["total_coin"]
        prow.append([n, h["total_invested_thb"], h["total_coin"], avg, info.get("price"), None, None,
                     info.get("signal"), (info.get("risk_pct") or 0) / 100 if info.get("risk_pct") is not None else None,
                     avg * (1 + tp), avg * (1 - cz), avg * (1 - sl), h.get("buy_count"), h.get("last_bought_date")])
    dtp, dsl = getattr(m, "DIP_TP_PCT", 6.0) / 100, getattr(m, "DIP_SL_PCT", 20.0) / 100
    for n, d in dips.items():
        if d.get("coin", 0) <= 0:
            continue
        info = latest.get(n) or {}
        avg = d["invested_thb"] / d["coin"]
        prow.append([f"{n} (ช้อน)", d["invested_thb"], d["coin"], avg, info.get("price") or d.get("current_price"), None, None,
                     f"ร่วง {d.get('drop_pct')}% ตอนช้อน", (info.get("risk_pct") or 0) / 100 if info.get("risk_pct") is not None else None,
                     avg * (1 + dtp), None, avg * (1 - dsl), 1, (d.get("bought_at") or "")[:10]])
    end = table(ws, 3, ["เหรียญ", "ทุน", "จำนวนเหรียญ", "ทุนเฉลี่ย/เหรียญ", "ราคาล่าสุด", "มูลค่า", "กำไร/ขาดทุน",
                        "แนวโน้ม", "ผันผวน 48 ชม.", f"ราคา TP (+{tp:.0%})", f"ราคาเริ่มตัด DC (-{cz:.0%})",
                        f"ราคา SL (-{sl:.0%})", "ซื้อกี่ครั้ง", "ซื้อล่าสุด"], prow,
                widths=[9, 11, 16, 16, 15, 11, 12, 24, 12, 15, 17, 15, 12, 13],
                formats={2: THB, 3: "0.00000000", 4: "#,##0.00####", 5: "#,##0.00####", 6: THB, 7: PCT,
                         9: "0.00%", 10: "#,##0.00####", 11: "#,##0.00####", 12: "#,##0.00####"})
    for r in range(4, end + 1):
        ws.cell(row=r, column=6, value=f"=IF(E{r}=\"\",\"\",C{r}*E{r})").number_format = THB
        ws.cell(row=r, column=7, value=f"=IF(E{r}=\"\",\"\",F{r}/B{r}-1)").number_format = PCT
    if prow:
        ws.cell(row=end + 1, column=1, value="รวม").font = B
        ws.cell(row=end + 1, column=2, value=f"=SUM(B4:B{end})").number_format = THB
        ws.cell(row=end + 1, column=6, value=f"=SUM(F4:F{end})").number_format = THB
        ws.cell(row=end + 1, column=7, value=f"=IF(B{end+1}=0,\"\",F{end+1}/B{end+1}-1)").number_format = PCT
        ws.cell(row=end + 2, column=1, value="เงินบาทในกอง").font = N
        ws.cell(row=end + 2, column=6, value=pool).number_format = THB

    # ---------------- ระบบตัดสินใจ ----------------
    ws = wb.create_sheet("ระบบตัดสินใจ")
    title(ws, "ระบบตัดสินใจ (ค่าที่บอทใช้อยู่จริงตอนนี้)")
    g = lambda k, d=None: getattr(m, k, d)
    rules = [
        ("ขาย", "TP (ขายทำกำไร)", f"+{g('TAKE_PROFIT_PCT')}% จากทุนเฉลี่ย → ขายทั้งหมด"),
        ("ขาย", "ตัดตาม Death Cross", (f"ขาดทุนเกิน {g('CUT_ZONE_PCT')}% และ MA12 ต่ำกว่า MA24 (กราฟรายชั่วโมง)"
                                       if g("SIGNAL_CUT_ENABLED") else "ปิดอยู่")),
        ("ขาย", "SL (ขายแน่นอน)", f"-{g('STOP_LOSS_PCT')}% จากทุนเฉลี่ย"),
        ("ซื้อ", "เลือกเหรียญ", "10 อันดับแรกตามมูลค่าซื้อขาย 24 ชม. บน Bitkub + BTC/ETH เสมอ · ตัดเหรียญที่ซื้อผ่าน API ไม่ได้"),
        ("ซื้อ", "ลำดับซื้อ", f"Golden Cross ที่ผันผวน ≤ {g('GC_PRIORITY_MAX_RISK_PCT')}% ได้ก่อน → ที่เหลือเรียงผันผวน 48 ชม. น้อยไปมาก"),
        ("ซื้อ", "ขนาดซื้อ", f"เติมแต่ละเหรียญให้ทุนครบ {g('DCA_UNIT_THB')} บาท · ทุกตัวครบแล้วซื้อซ้ำได้ถึง {g('MAX_PER_COIN_THB')} บาท · ขั้นต่ำ {g('MIN_ORDER_THB')} บาท"),
        ("ซื้อ", "ซื้อกลับหลังขาย", f"รอ {g('REBUY_COOLDOWN_HOURS')} ชม. + Golden Cross + ผันผวน ≤ {g('REBUY_MAX_RISK_PCT')}%"),
        ("จังหวะ", "รอบทำงาน", "จัดลำดับใหม่ทุกชั่วโมง · เช็คราคาและซื้อขายทุก 1 นาที"),
        ("เงิน", "กองเงิน", ("ทั้งบัญชี: เงินบาทในบัญชี = กองเงินบอท · กำไรทบลงทุนต่อ" if g("WHOLE_ACCOUNT_MODE")
                             else "กองเงินแยกจากบัญชี")),
        ("เงิน", "แจ้งถอน", f"พอร์ตโตเกินทุน {g('PROFIT_WITHDRAW_ALERT_THB')} บาท → แจ้ง Telegram (ถอนเองในแอป แล้วพิมพ์ \"ถอน X\")"),
        ("บันทึก", "ยังไม่ใช้ตัดสินใจ", "premium เทียบตลาดโลก · ความหนากระดาน · แรงซื้อ 200 รายการล่าสุด · ราคา 24 ชม. · แนวโน้ม BTC (รอพิสูจน์ ~30 ไม้)"),
    ]
    table(ws, 3, ["กลุ่ม", "กติกา", "รายละเอียด"], rules, widths=[10, 22, 100])
    hist = _jsonl("analysis/config_history.jsonl")
    r0 = 6 + len(rules)
    ws.cell(row=r0 - 1, column=1, value="ประวัติการเปลี่ยนกติกา").font = B
    table(ws, r0, ["เวลา", "TP %", "SL %", "ตัดตาม DC", "จุดเริ่มตัด %", "ทุน/เหรียญ", "ค่าอื่นๆ"],
          [[_bkk(h.get("at")), h["config"].get("take_profit_pct"), h["config"].get("stop_loss_pct"),
            "เปิด" if h["config"].get("signal_cut_enabled") else "ปิด", h["config"].get("cut_zone_pct"),
            h["config"].get("unit_thb"),
            json.dumps({k: v for k, v in h["config"].items() if k not in (
                "take_profit_pct", "stop_loss_pct", "signal_cut_enabled", "cut_zone_pct", "unit_thb")},
                ensure_ascii=False)] for h in hist if isinstance(h.get("config"), dict)])
    ws.freeze_panes = None

    # ---------------- ประวัติซื้อ (journal: รายไม้ พร้อมผล) ----------------
    ws = wb.create_sheet("ประวัติซื้อ")
    title(ws, "ทุกไม้ที่ซื้อ และผลลัพธ์ (trade journal)", "สภาพตลาดตอนซื้อบันทึกไว้เพื่อหาว่าปัจจัยไหนช่วยให้ชนะ")
    lrows = []
    for l in sorted(journal.get("lots", []), key=lambda x: x.get("opened_at") or ""):
        f = ((l.get("features") or {}).get("coin") or {})
        lrows.append([_bkk(l.get("opened_at")), l.get("coin"), l.get("invested_thb"), l.get("fill_price"),
                      "ปิดแล้ว" if l.get("status") == "closed" else "ถืออยู่", _bkk(l.get("closed_at")),
                      l.get("exit_reason"), l.get("exit_price"), l.get("proceeds_thb"), l.get("pnl_thb"),
                      (l["pnl_pct"] / 100) if l.get("pnl_pct") is not None else None,
                      f.get("signal"), f.get("risk_pct"), f.get("premium_pct"), f.get("depth_imbalance"),
                      f.get("buy_flow_pct"), "ใช่" if l.get("gc_priority") else "", "ใช่" if l.get("rebuy") else "",
                      l.get("note") or ""])
    table(ws, 3, ["ซื้อเมื่อ", "เหรียญ", "ทุน", "ราคาซื้อ", "สถานะ", "ปิดเมื่อ", "เหตุผลขาย", "ราคาขาย", "ได้คืน",
                  "กำไร/ขาดทุน", "%", "แนวโน้มตอนซื้อ", "ผันผวน %", "premium %", "ความหนากระดาน",
                  "แรงซื้อ %", "GC ได้ก่อน", "ซื้อกลับ", "หมายเหตุ"], lrows,
          widths=[16, 8, 10, 14, 9, 16, 12, 14, 10, 11, 9, 22, 9, 10, 12, 9, 9, 8, 30],
          formats={3: THB, 4: "#,##0.00####", 8: "#,##0.00####", 9: THB, 10: THB, 11: PCT})
    for r in range(4, 4 + len(lrows)):
        v = ws.cell(row=r, column=10).value
        if isinstance(v, (int, float)):
            ws.cell(row=r, column=10).font = GOOD if v > 0 else BAD

    # ---------------- ประวัติขาย ----------------
    ws = wb.create_sheet("ประวัติขาย")
    title(ws, "ทุกครั้งที่ขาย")
    reason_th = {"take_profit": "ขายทำกำไร (TP)", "stop_loss": "ตัดขาดทุน (SL)", "signal_cut": "ตัดตาม Death Cross",
                 "rebalance": "ปรับสัดส่วน", "dip_take_profit": "กองช้อน: ขายทำกำไร", "dip_stop_loss": "กองช้อน: ตัดขาดทุน"}
    srows = [[s.get("date"), s.get("name"), reason_th.get(s.get("reason"), s.get("reason")), s.get("sold_coin"),
              s.get("sell_price"), s.get("invested_thb"), s.get("proceeds_thb"), s.get("realized_pnl_thb"),
              (s["realized_pnl_thb"] / s["invested_thb"]) if s.get("invested_thb") else None] for s in sells]
    end = table(ws, 3, ["วันที่", "เหรียญ", "เหตุผล", "จำนวนเหรียญ", "ราคาขาย", "ทุน", "ได้คืน", "กำไร/ขาดทุน", "%"],
                srows, widths=[12, 8, 20, 16, 14, 10, 10, 12, 9],
                formats={4: "0.00000000", 5: "#,##0.00####", 6: THB, 7: THB, 8: THB, 9: PCT})
    for r in range(4, end + 1):
        v = ws.cell(row=r, column=8).value
        if isinstance(v, (int, float)):
            ws.cell(row=r, column=8).font = GOOD if v > 0 else BAD
    if srows:
        ws.cell(row=end + 1, column=1, value="รวม").font = B
        for c, col in ((6, "F"), (7, "G"), (8, "H")):
            ws.cell(row=end + 1, column=c, value=f"=SUM({col}4:{col}{end})").number_format = THB

    # ---------------- กำไรขาดทุน ----------------
    ws = wb.create_sheet("กำไรขาดทุน")
    title(ws, "กำไร/ขาดทุนที่ขายแล้ว แยกตามเหรียญ / เหตุผล / วัน")
    by_coin = defaultdict(lambda: [0, 0, 0.0, 0.0])
    by_reason = defaultdict(lambda: [0, 0.0])
    by_day = OrderedDict()
    for s in sells:
        pnl = s.get("realized_pnl_thb") or 0.0
        bc = by_coin[s.get("name")]
        bc[0] += 1
        bc[1] += 1 if pnl > 0 else 0
        bc[2] += s.get("invested_thb") or 0.0
        bc[3] += pnl
        br = by_reason[reason_th.get(s.get("reason"), s.get("reason"))]
        br[0] += 1
        br[1] += pnl
        by_day[s.get("date")] = by_day.get(s.get("date"), 0.0) + pnl
    crow = [[k, v[0], v[1], round(v[2], 2), round(v[3], 2)] for k, v in sorted(by_coin.items(), key=lambda kv: -kv[1][3])]
    end = table(ws, 3, ["เหรียญ", "ขายกี่ครั้ง", "ชนะ", "ทุนรวมที่ขาย", "กำไร/ขาดทุนรวม"], crow,
                widths=[18, 12, 8, 14, 16, 4, 18, 12, 16], formats={4: THB, 5: THB})
    if crow:
        ch = BarChart()
        ch.title, ch.y_axis.title = "กำไร/ขาดทุนรวมรายเหรียญ (บาท)", "บาท"
        ch.add_data(Reference(ws, min_col=5, min_row=3, max_row=end), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=end))
        ch.height, ch.width, ch.legend = 8, 16, None
        ch.x_axis.tickLblPos = "low"
        ws.add_chart(ch, f"A{end + 3}")
    # เหตุผล + รายวัน ด้านขวา
    for c, h in enumerate(["เหตุผลขาย", "ครั้ง", "กำไร/ขาดทุนรวม"], 7):
        x = ws.cell(row=3, column=c, value=h)
        x.font, x.fill, x.border = H, HF, BOX
    for i, (k, v) in enumerate(by_reason.items(), 4):
        ws.cell(row=i, column=7, value=k).font = N
        ws.cell(row=i, column=8, value=v[0]).font = N
        ws.cell(row=i, column=9, value=round(v[1], 2)).number_format = THB
    d0 = 6 + len(by_reason)
    for c, h in enumerate(["วันที่", "กำไร/ขาดทุนของวัน", "สะสม"], 7):
        x = ws.cell(row=d0, column=c, value=h)
        x.font, x.fill, x.border = H, HF, BOX
    for i, (k, v) in enumerate(by_day.items(), d0 + 1):
        ws.cell(row=i, column=7, value=k).font = N
        ws.cell(row=i, column=8, value=round(v, 2)).number_format = THB
        ws.cell(row=i, column=9, value=f"=SUM(H${d0 + 1}:H{i})").number_format = THB
    if by_day:
        ch = LineChart()
        ch.title, ch.y_axis.title = "กำไร/ขาดทุนที่ขายแล้ว สะสมรายวัน (บาท)", "บาท"
        ch.add_data(Reference(ws, min_col=9, min_row=d0, max_row=d0 + len(by_day)), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=7, min_row=d0 + 1, max_row=d0 + len(by_day)))
        ch.height, ch.width, ch.legend = 8, 16, None
        for sr in ch.series:
            sr.smooth = False
        ws.add_chart(ch, f"K{d0}")
    ws.column_dimensions["G"].width = 20
    ws.column_dimensions["I"].width = 14

    # ---------------- มูลค่าพอร์ต ----------------
    ws = wb.create_sheet("มูลค่าพอร์ต")
    title(ws, "มูลค่าพอร์ตรายชั่วโมง",
          "ก่อนเริ่มโหมดทั้งบัญชี = เฉพาะเงินของบอท (ทุน = เงินที่เติมเข้าบอท) · หลังจากนั้น = ทั้งบัญชี (มูลค่ากระโดดตอนเปลี่ยนโหมด ไม่ใช่กำไร — ดูกราฟกำไรเทียบทุนแทน)")
    eq = [e for e in _jsonl("analysis/backfill/equity.jsonl") + _jsonl("analysis/equity.jsonl")
          if e.get("total_equity_thb") is not None and not e.get("invalid")]
    eq.sort(key=lambda e: e.get("at") or "")
    erows = []
    for e in eq:
        cap = _capital_of(e)
        erows.append([_bkk(e.get("at")), e.get("total_equity_thb"), cap,
                      round(e["total_equity_thb"] - cap, 2), e.get("holdings_value_thb"), e.get("cash_pool_thb"),
                      e.get("profit_reserved_thb"), e.get("realized_pnl_total_thb")])
    end = table(ws, 3, ["เวลา", "มูลค่ารวม", "ทุน", "กำไร/ขาดทุนเทียบทุน", "มูลค่าเหรียญ", "เงินในกอง",
                        "กำไรแยกเก็บ (เดิม)", "กำไรที่ขายแล้วสะสม"], erows,
                widths=[17, 12, 12, 18, 13, 12, 16, 18],
                formats={c: THB for c in range(2, 9)})
    if len(erows) > 1:
        ch = LineChart()
        ch.title, ch.y_axis.title = "กำไร/ขาดทุนเทียบทุน (บาท)", "บาท"
        ch.add_data(Reference(ws, min_col=4, min_row=3, max_row=end), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=end))
        ch.x_axis.number_format, ch.height, ch.width, ch.legend = "dd/mm hh:mm", 9, 22, None
        ch.x_axis.tickLblPos = "low"
        for sr in ch.series:
            sr.smooth = False
        ws.add_chart(ch, "J3")
        ch2 = LineChart()
        ch2.title, ch2.y_axis.title = "มูลค่าพอร์ต vs ทุน (บาท)", "บาท"
        ch2.add_data(Reference(ws, min_col=2, max_col=3, min_row=3, max_row=end), titles_from_data=True)
        ch2.set_categories(Reference(ws, min_col=1, min_row=4, max_row=end))
        ch2.x_axis.number_format, ch2.height, ch2.width = "dd/mm hh:mm", 9, 22
        for sr in ch2.series:
            sr.smooth = False
        ws.add_chart(ch2, "J23")

    # ---------------- เหตุการณ์ (เหตุผลการตัดสินใจ) ----------------
    ws = wb.create_sheet("เหตุการณ์")
    title(ws, "เหตุผลการตัดสินใจของบอท (ล่าสุด 3,000 รายการ)")
    kind_th = {"buy": "ซื้อ", "sell": "ขาย", "buy_failed": "ซื้อไม่สำเร็จ", "rebuy_skipped": "ยังไม่ซื้อกลับ",
               "hold_in_loss": "ถือรอ (ติดลบ)", "withdraw": "ถอน", "whole_account_start": "เริ่มโหมดทั้งบัญชี",
               "holding_gone": "เหรียญหายจากบัญชี"}
    evs = []
    for fpath in sorted(glob.glob("analysis/events/*.jsonl")):
        evs += _jsonl(fpath)
    evs = [e for e in evs if e.get("kind") != "hold_in_loss"][-3000:]
    vrows = []
    for e in reversed(evs):
        info = {k: v for k, v in e.items() if k not in ("at", "kind", "coin")}
        vrows.append([_bkk(e.get("at")), kind_th.get(e.get("kind"), e.get("kind")), e.get("coin"),
                      info.get("reason") or "", info.get("thb") or info.get("proceeds_thb"),
                      json.dumps(info, ensure_ascii=False)])
    table(ws, 3, ["เวลา", "เหตุการณ์", "เหรียญ", "เหตุผล", "บาท", "รายละเอียด"], vrows,
          widths=[17, 18, 8, 34, 10, 90], formats={5: THB})

    try:
        wb.calculation.fullCalcOnLoad = True
    except AttributeError:
        pass
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)
    return path


# ======================= CSV สำหรับ Google Sheets (IMPORTDATA) =======================
CSV_DIR = "reports/csv"
REASON_TH = {"take_profit": "ขายทำกำไร (TP)", "stop_loss": "ตัดขาดทุน (SL)", "signal_cut": "ตัดตาม Death Cross",
             "rebalance": "ปรับสัดส่วน", "dip_take_profit": "กองช้อน: ขายทำกำไร", "dip_stop_loss": "กองช้อน: ตัดขาดทุน"}
KIND_TH = {"buy": "ซื้อ", "sell": "ขาย", "buy_failed": "ซื้อไม่สำเร็จ", "rebuy_skipped": "ยังไม่ซื้อกลับ",
           "hold_in_loss": "ถือรอ (ติดลบ)", "withdraw": "ถอน", "whole_account_start": "เริ่มโหมดทั้งบัญชี",
           "holding_gone": "เหรียญหายจากบัญชี"}


def _t(dt):
    return dt.strftime("%Y-%m-%d %H:%M") if isinstance(dt, datetime) else ("" if dt is None else dt)


def _r(x, n=2):
    return "" if x is None else round(x, n)


def export_csv(m=None, out_dir=CSV_DIR):
    """เขียนข้อมูลแต่ละแท็บเป็น CSV (ค่าคำนวณแล้ว ไม่มีสูตร) ให้ Google Sheets ดึงด้วย IMPORTDATA"""
    import csv
    state = _load("state/dca_state.json", {})
    journal = _load("state/trade_journal.json", {"lots": []})
    latest = _load("latest.json", {})
    g = lambda k, d=None: getattr(m, k, d)
    now = datetime.now(BKK).replace(tzinfo=None)
    tables = {}

    holdings = {n: h for n, h in (state.get("holdings") or {}).items() if (h or {}).get("total_coin", 0) > 0}
    tp, sl, cz = (g("TAKE_PROFIT_PCT", 6.0) or 0) / 100, (g("STOP_LOSS_PCT", 20.0) or 0) / 100, (g("CUT_ZONE_PCT", 10.0) or 0) / 100
    hrows, value_now = [], 0.0
    for n, h in sorted(holdings.items(), key=lambda kv: -kv[1]["total_invested_thb"]):
        info = latest.get(n) or {}
        avg = h["total_invested_thb"] / h["total_coin"]
        px = info.get("price")
        val = h["total_coin"] * px if px else None
        value_now += val or 0.0
        hrows.append([n, _r(h["total_invested_thb"]), h["total_coin"], _r(avg, 6), px, _r(val),
                      _r(val - h["total_invested_thb"]) if val is not None else "",
                      _r((val / h["total_invested_thb"] - 1) * 100) if val is not None else "",
                      info.get("signal") or "", info.get("risk_pct"), _r(avg * (1 + tp), 6), _r(avg * (1 - cz), 6),
                      _r(avg * (1 - sl), 6), h.get("buy_count"), h.get("last_bought_date") or ""])
    dtp, dsl = (g("DIP_TP_PCT", 6.0) or 0) / 100, (g("DIP_SL_PCT", 20.0) or 0) / 100
    for n, d in (state.get("dip_positions") or {}).items():
        if d.get("coin", 0) <= 0:
            continue
        info = latest.get(n) or {}
        avg = d["invested_thb"] / d["coin"]
        px = info.get("price") or d.get("current_price")
        val = d["coin"] * px if px else None
        value_now += val or 0.0
        hrows.append([f"{n} (ช้อน)", _r(d["invested_thb"]), d["coin"], _r(avg, 6), px, _r(val),
                      _r(val - d["invested_thb"]) if val is not None else "",
                      _r((val / d["invested_thb"] - 1) * 100) if val is not None else "",
                      f"ร่วง {d.get('drop_pct')}% ตอนช้อน", info.get("risk_pct"), _r(avg * (1 + dtp), 6), "",
                      _r(avg * (1 - dsl), 6), 1, (d.get("bought_at") or "")[:10]])
    tables["holdings"] = (["เหรียญ", "ทุน (บาท)", "จำนวนเหรียญ", "ทุนเฉลี่ย/เหรียญ", "ราคาล่าสุด", "มูลค่า (บาท)",
                           "กำไร/ขาดทุน (บาท)", "กำไร/ขาดทุน (%)", "แนวโน้ม", "ผันผวน 48 ชม. (%)",
                           f"ราคา TP (+{tp*100:.0f}%)", f"ราคาเริ่มตัด DC (-{cz*100:.0f}%)", f"ราคา SL (-{sl*100:.0f}%)",
                           "ซื้อกี่ครั้ง", "ซื้อล่าสุด"], hrows)

    pool = state.get("cash_pool_thb", 0.0)
    capital = state.get("capital_thb")
    equity = value_now + pool
    closed = [l for l in journal.get("lots", []) if l.get("status") == "closed" and l.get("pnl_pct") is not None]
    wins = [l for l in closed if l.get("win")]
    sells = state.get("sell_log") or []
    since = state.get("whole_account_since")
    closed_new = [l for l in closed if since and (l.get("closed_at") or "") >= since]
    tables["overview"] = (["รายการ", "ค่า"], [
        ["อัปเดตล่าสุด (เวลาไทย)", _t(now)],
        ["ทุนทั้งบัญชี (บาท)", _r(capital)],
        ["มูลค่าเหรียญตอนนี้ (บาท)", _r(value_now)],
        ["เงินบาทในกอง (บาท)", _r(pool)],
        ["มูลค่าทั้งบัญชี (บาท)", _r(equity)],
        ["กำไร/ขาดทุนเทียบทุน (บาท)", _r(equity - capital) if capital else ""],
        ["กำไร/ขาดทุนเทียบทุน (%)", _r((equity / capital - 1) * 100) if capital else ""],
        ["แจ้งถอนเมื่อกำไรถึง (บาท)", g("PROFIT_WITHDRAW_ALERT_THB", 300)],
        ["กำไรที่ขายแล้วสะสม ตั้งแต่เริ่มบอท (บาท)", _r(state.get("realized_pnl_total_thb", 0.0))],
        ["ถอนออกสะสม (บาท)", _r(state.get("withdrawn_total_thb", 0.0))],
        ["ไม้ที่ปิดแล้ว (ทั้งหมด)", len(closed)],
        ["Winrate ทั้งหมด (%)", _r(len(wins) / len(closed) * 100, 1) if closed else ""],
        ["ไม้ที่ปิดแล้ว (กติกาปัจจุบัน ตั้งแต่โหมดทั้งบัญชี)", len(closed_new)],
        ["Winrate กติกาปัจจุบัน (%)", _r(sum(1 for l in closed_new if l.get("win")) / len(closed_new) * 100, 1) if closed_new else ""],
        ["จำนวนครั้งที่ขาย (ทั้งหมด)", len(sells)],
        ["เหรียญที่ถืออยู่", len(holdings)],
        ["กองช้อน: เป้า (บาท)", _r((capital or 0) * (g("DIP_RESERVE_PCT", 0) or 0) / 100)],
        ["กองช้อน: เงินรอช้อน (บาท)", _r(g("dip_reserve_cash")(state)) if g("dip_reserve_cash") else ""],
        ["กองช้อน: ไม้ที่เปิดอยู่", len(state.get("dip_positions") or {})],
        ["TP / ตัดตาม DC / SL (%)", f"+{g('TAKE_PROFIT_PCT')} / -{g('CUT_ZONE_PCT')} (เมื่อ Death Cross) / -{g('STOP_LOSS_PCT')}"],
    ])

    tables["rules"] = (["กลุ่ม", "กติกา", "รายละเอียด"], [
        ["ขาย", "TP (ขายทำกำไร)", f"+{g('TAKE_PROFIT_PCT')}% จากทุนเฉลี่ย → ขายทั้งหมด"],
        ["ขาย", "ตัดตาม Death Cross", (f"ขาดทุนเกิน {g('CUT_ZONE_PCT')}% และ MA12 ต่ำกว่า MA24 (กราฟรายชั่วโมง)"
                                       if g("SIGNAL_CUT_ENABLED") else "ปิดอยู่")],
        ["ขาย", "SL (ขายแน่นอน)", f"-{g('STOP_LOSS_PCT')}% จากทุนเฉลี่ย"],
        ["ซื้อ", "เลือกเหรียญ", "10 อันดับแรกตามมูลค่าซื้อขาย 24 ชม. บน Bitkub + BTC/ETH เสมอ · ตัดเหรียญที่ซื้อผ่าน API ไม่ได้"],
        ["ซื้อ", "ลำดับซื้อ", f"Golden Cross ที่ผันผวน ≤ {g('GC_PRIORITY_MAX_RISK_PCT')}% ได้ก่อน → ที่เหลือเรียงผันผวน 48 ชม. น้อยไปมาก"],
        ["ซื้อ", "ขนาดซื้อ", f"เติมแต่ละเหรียญให้ทุนครบ {g('DCA_UNIT_THB')} บาท · ครบทุกตัวแล้วซื้อซ้ำได้ถึง {g('MAX_PER_COIN_THB')} บาท · ขั้นต่ำ {g('MIN_ORDER_THB')} บาท"],
        ["ซื้อ", "ซื้อกลับหลังขาย", f"รอ {g('REBUY_COOLDOWN_HOURS')} ชม. + Golden Cross + ผันผวน ≤ {g('REBUY_MAX_RISK_PCT')}%"],
        ["จังหวะ", "รอบทำงาน", "จัดลำดับใหม่ทุกชั่วโมง · เช็คราคาและซื้อขายทุก 1 นาที"],
        ["เงิน", "กองเงิน", "ทั้งบัญชี: เงินบาทในบัญชี = กองเงินบอท · กำไรทบลงทุนต่อ" if g("WHOLE_ACCOUNT_MODE") else "กองเงินแยก"],
        ["เงิน", "แจ้งถอน", f"พอร์ตโตเกินทุน {g('PROFIT_WITHDRAW_ALERT_THB')} บาท → แจ้ง Telegram (ถอนเองในแอป แล้วพิมพ์ \"ถอน X\")"],
        ["กองช้อน", "ขนาดกอง", f"{g('DIP_RESERVE_PCT')}% ของทุนทั้งบัญชี · เติมจากเงินที่ได้คืนตอนขาย (เงินส่วนนี้ห้ามใช้ซื้อแบบปกติ)"],
        ["กองช้อน", "ช้อนซื้อ", f"ราคาต่ำกว่าจุดสูงสุด 72 ชม. ตั้งแต่ {g('DIP_DROP_PCT')}% และผันผวน 48 ชม. ≤ {g('DIP_MAX_RISK_PCT')}% → ซื้อ {g('DIP_UNIT_THB')} บาท (ร่วงลึกสุดก่อน)"],
        ["กองช้อน", "ขายไม้ช้อน", f"+{g('DIP_TP_PCT')}% ขายทำกำไร · -{g('DIP_SL_PCT')}% ตัดขาดทุน · ขายแล้วรอ {g('DIP_COOLDOWN_HOURS')} ชม. ก่อนช้อนเหรียญเดิม"],
        ["ระบบปรับเอง", "ทดลองเงียบ → ปล่อยใช้จริง", "ทุกวันอาทิตย์ 08:41 น. หาค่าที่ดีกว่าจากกราฟ Bitkub 150 วัน (ต้องดีกว่า ≥10% ทั้งช่วงหาค่าและช่วงทดสอบ) · ทดลองเงียบเทียบค่าจริงทุกสัปดาห์ · ชนะติดกัน 3 สัปดาห์ถึงปล่อยใช้จริง · กรอบ TP 4-10% / SL 15-25% / ตัด DC 8-15% / ช้อน 10-20% / ขายไม้ช้อน 4-10%"
         + (f" · ใช้อยู่: {g('TUNED_APPLIED')}" if g("TUNED_APPLIED") else "")],
        ["บันทึก", "ยังไม่ใช้ตัดสินใจ", "premium · ความหนากระดาน · แรงซื้อ 200 รายการล่าสุด · ราคา 24 ชม. · แนวโน้ม BTC (รอพิสูจน์ ~30 ไม้)"],
    ])

    brow = []
    for l in sorted(journal.get("lots", []), key=lambda x: x.get("opened_at") or "", reverse=True):
        f = ((l.get("features") or {}).get("coin") or {})
        brow.append([_t(_bkk(l.get("opened_at"))), l.get("coin"), l.get("invested_thb"), l.get("fill_price"),
                     "ปิดแล้ว" if l.get("status") == "closed" else "ถืออยู่", _t(_bkk(l.get("closed_at"))),
                     REASON_TH.get(l.get("exit_reason"), l.get("exit_reason") or ""), l.get("exit_price") or "",
                     _r(l.get("proceeds_thb")), _r(l.get("pnl_thb")), _r(l.get("pnl_pct")), f.get("signal") or "",
                     f.get("risk_pct") or "", f.get("premium_pct") or "", f.get("depth_imbalance") or "",
                     f.get("buy_flow_pct") or "", "ใช่" if l.get("gc_priority") else "", "ใช่" if l.get("rebuy") else "",
                     l.get("note") or ""])
    tables["buys"] = (["ซื้อเมื่อ", "เหรียญ", "ทุน (บาท)", "ราคาซื้อ", "สถานะ", "ปิดเมื่อ", "เหตุผลขาย", "ราคาขาย",
                       "ได้คืน (บาท)", "กำไร/ขาดทุน (บาท)", "กำไร/ขาดทุน (%)", "แนวโน้มตอนซื้อ", "ผันผวน (%)",
                       "premium (%)", "ความหนากระดาน", "แรงซื้อ (%)", "GC ได้ก่อน", "ซื้อกลับ", "หมายเหตุ"], brow)

    tables["sells"] = (["วันที่", "เหรียญ", "เหตุผล", "จำนวนเหรียญ", "ราคาขาย", "ทุน (บาท)", "ได้คืน (บาท)",
                        "กำไร/ขาดทุน (บาท)", "กำไร/ขาดทุน (%)"],
                       [[s.get("date"), s.get("name"), REASON_TH.get(s.get("reason"), s.get("reason")), s.get("sold_coin"),
                         s.get("sell_price"), s.get("invested_thb"), s.get("proceeds_thb"), s.get("realized_pnl_thb"),
                         _r(s["realized_pnl_thb"] / s["invested_thb"] * 100) if s.get("invested_thb") else ""]
                        for s in reversed(sells)])

    by_coin = defaultdict(lambda: [0, 0, 0.0, 0.0])
    by_day = OrderedDict()
    for s in sells:
        pnl = s.get("realized_pnl_thb") or 0.0
        bc = by_coin[s.get("name")]
        bc[0] += 1; bc[1] += 1 if pnl > 0 else 0; bc[2] += s.get("invested_thb") or 0.0; bc[3] += pnl
        by_day[s.get("date")] = by_day.get(s.get("date"), 0.0) + pnl
    tables["pnl_coin"] = (["เหรียญ", "ขายกี่ครั้ง", "ชนะ", "ทุนรวมที่ขาย (บาท)", "กำไร/ขาดทุนรวม (บาท)"],
                          [[k, v[0], v[1], _r(v[2]), _r(v[3])] for k, v in sorted(by_coin.items(), key=lambda kv: -kv[1][3])])
    cum, drows = 0.0, []
    for k, v in by_day.items():
        cum += v
        drows.append([k, _r(v), _r(cum)])
    tables["pnl_day"] = (["วันที่", "กำไร/ขาดทุนของวัน (บาท)", "สะสม (บาท)"], drows)

    eq = [e for e in _jsonl("analysis/backfill/equity.jsonl") + _jsonl("analysis/equity.jsonl")
          if e.get("total_equity_thb") is not None and not e.get("invalid")]
    eq.sort(key=lambda e: e.get("at") or "")
    erows = []
    for e in eq:
        cap = _capital_of(e)
        erows.append([_t(_bkk(e.get("at"))), e["total_equity_thb"], cap, _r(e["total_equity_thb"] - cap),
                      e.get("holdings_value_thb"), e.get("cash_pool_thb")])
    tables["equity"] = (["เวลา", "มูลค่ารวม (บาท)", "ทุน (บาท)", "กำไร/ขาดทุนเทียบทุน (บาท)", "มูลค่าเหรียญ (บาท)",
                         "เงินในกอง (บาท)"], erows)

    evs = []
    for fpath in sorted(glob.glob("analysis/events/*.jsonl")):
        evs += _jsonl(fpath)
    evs = [e for e in evs if e.get("kind") != "hold_in_loss"][-1000:]
    tables["events"] = (["เวลา", "เหตุการณ์", "เหรียญ", "เหตุผล", "บาท", "รายละเอียด"],
                        [[_t(_bkk(e.get("at"))), KIND_TH.get(e.get("kind"), e.get("kind")), e.get("coin") or "",
                          e.get("reason") or "", e.get("thb") or e.get("proceeds_thb") or "",
                          json.dumps({k: v for k, v in e.items() if k not in ("at", "kind", "coin")}, ensure_ascii=False)]
                         for e in reversed(evs)])

    orows = []
    for r in reversed(_jsonl("analysis/optimizer/history.jsonl")):
        live, cand = r.get("live") or {}, r.get("candidate") or {}
        diff = ", ".join(f"{k} {live.get(k)}→{cand.get(k)}" for k in cand if cand.get(k) != live.get(k)) or "ไม่เปลี่ยน"
        ev = r.get("evidence") or {}
        orows.append([_t(_bkk(r.get("at"))), "ปล่อยใช้จริง" if r.get("promoted") else "ทดลองเงียบ", diff, r.get("streak"),
                      ((ev.get("main") or {}).get("live_test") or {}).get("per_day"), ((ev.get("main") or {}).get("test") or {}).get("per_day"),
                      ((ev.get("dip") or {}).get("live_test") or {}).get("per_day"), ((ev.get("dip") or {}).get("test") or {}).get("per_day")])
    tables["optimizer"] = (["เวลา", "สถานะ", "ค่าที่ทดลอง (ค่าจริง→ค่าทดลอง)", "ชนะติดกัน (สัปดาห์)",
                            "กองปกติ ค่าจริง %/วัน", "กองปกติ ค่าทดลอง %/วัน", "กองช้อน ค่าจริง %/วัน", "กองช้อน ค่าทดลอง %/วัน"], orows)

    os.makedirs(out_dir, exist_ok=True)
    for name, (headers, rows) in tables.items():
        tmp = f"{out_dir}/{name}.csv.tmp"
        with open(tmp, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(headers)
            w.writerows(rows)
        os.replace(tmp, f"{out_dir}/{name}.csv")
    return sorted(tables)


def safe_build(dca_module=None):
    path = None
    try:
        path = build(dca_module)
    except Exception as e:  # noqa: BLE001 — รายงานพังต้องไม่ทำให้บอทหยุดเทรด
        print(f"สร้างไฟล์ Excel ไม่สำเร็จ: {e}")
    try:
        export_csv(dca_module)
    except Exception as e:  # noqa: BLE001
        print(f"สร้างไฟล์ CSV ไม่สำเร็จ: {e}")
    return path


if __name__ == "__main__":
    import dca_bot  # noqa: F401
    print(safe_build(dca_bot))
